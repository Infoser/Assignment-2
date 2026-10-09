"""Seed the GHG platform database from the GHG Sheet workbook.

Creates:
  - Versioned EmissionFactors: a 2023 revision (expired) and the current 2024
    revision for every activity, so historical accuracy is demonstrable.
  - EmissionRecords for 2023 and 2024. Quarterly Scope 1 & 2 quantities are
    split evenly across the months of their quarter to give the dashboard a
    monthly trend line. All values are converted to the canonical unit
    kgCO2e (sheet factors are tCO2/unit, so x1000).
  - BusinessMetrics ("Tons of Steel Produced", "Number of Employees").
  - One sample manual override with a matching AuditLog entry.

All synthetic variation is deterministic (md5-derived) so re-seeding is
reproducible.
"""

import hashlib
from datetime import date, timedelta

from openpyxl import load_workbook

from .database import DATA_XLSX, Base, engine, SessionLocal
from .models import AuditLog, BusinessMetric, EmissionFactor, EmissionRecord

YEAR = 2024
HIST_YEAR = 2023
SCOPE1_QUARTER_MONTHS = {"Q1": (1, 2, 3), "Q2": (4, 5, 6)}
SCOPE2_QUARTER_MONTHS = {"Q1": (1, 2, 3), "Q2": (4, 5, 6), "Q3": (7, 8, 9)}


def month_end(year: int, month: int) -> date:
    return (date(year, month, 28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)


def jitter(key: str, lo: float, hi: float) -> float:
    h = int(hashlib.md5(key.encode()).hexdigest(), 16)
    return lo + (h % 10000) / 10000.0 * (hi - lo)


class FactorRegistry:
    """Ensures exactly one factor version-set per (scope, activity, unit)."""

    def __init__(self, session):
        self.session = session
        self.cache: dict[tuple, dict[int, EmissionFactor]] = {}

    def get(self, scope: int, activity: str, unit: str, ef_tco2: float, source: str) -> dict[int, EmissionFactor]:
        key = (scope, activity, unit)
        if key in self.cache:
            return self.cache[key]
        ef_kgco2e = ef_tco2 * 1000.0
        hist_ratio = jitter(f"{activity}|{unit}|hist", 0.93, 0.97)
        versions = {
            HIST_YEAR: EmissionFactor(
                scope=scope, activity=activity, unit=unit,
                co2e_per_unit=round(ef_kgco2e * hist_ratio, 4), source=source,
                valid_from=date(HIST_YEAR, 1, 1), valid_to=date(HIST_YEAR, 12, 31), version=1,
            ),
            YEAR: EmissionFactor(
                scope=scope, activity=activity, unit=unit,
                co2e_per_unit=round(ef_kgco2e, 4), source=source,
                valid_from=date(YEAR, 1, 1), valid_to=None, version=2,
            ),
        }
        self.session.add_all(versions.values())
        self.session.flush()
        self.cache[key] = versions
        return versions


def add_record(session, scope, activity, section, plant, quantity, unit, activity_date, factor):
    emissions = quantity * factor.co2e_per_unit
    session.add(EmissionRecord(
        scope=scope, activity=activity, section=section, plant=plant,
        quantity=round(quantity, 4), unit=unit, activity_date=activity_date,
        factor_id=factor.id, factor_value_used=factor.co2e_per_unit,
        emissions_kgco2e=emissions,
    ))


def seed_scope1(session, registry):
    ws = load_workbook(DATA_XLSX, read_only=True, data_only=True)["Scope 1"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        section, material, unit, quarter, qty, ef_tco2, _ef_unit, _ghg, _cc, _ceff, source, plant = row[:12]
        if not material or qty is None:
            continue
        versions = registry.get(1, str(material).strip(), str(unit).strip(), float(ef_tco2), str(source).strip())
        months = SCOPE1_QUARTER_MONTHS[str(quarter).strip()]
        monthly_qty = float(qty) / len(months)
        for y in (YEAR, HIST_YEAR):
            factor = versions[y]
            for m in months:
                hist_key = f"s1|{material}|{quarter}|{m}"
                q = monthly_qty * (jitter(hist_key, 0.85, 0.95) if y == HIST_YEAR else 1.0)
                add_record(session, 1, factor.activity, str(section).strip(), str(plant).strip(),
                           q, factor.unit, month_end(y, m), factor)


def seed_scope2(session, registry):
    ws = load_workbook(DATA_XLSX, read_only=True, data_only=True)["Scope 2"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        section, energy_type, supplier, unit, consumed, ef_tco2, _emissions, _ren, _freq, _notes, _tou, source, _lf, _pf, plant, _unc, _int, quarter = row[:18]
        if not energy_type or consumed is None:
            continue
        activity = f"{str(energy_type).strip()} - {str(supplier).strip()}"
        versions = registry.get(2, activity, str(unit).strip(), float(ef_tco2), str(source).strip())
        months = SCOPE2_QUARTER_MONTHS[str(quarter).strip()]
        monthly_qty = float(consumed) / len(months)
        for y in (YEAR, HIST_YEAR):
            factor = versions[y]
            for m in months:
                hist_key = f"s2|{activity}|{quarter}|{m}"
                q = monthly_qty * (jitter(hist_key, 0.85, 0.95) if y == HIST_YEAR else 1.0)
                add_record(session, 2, factor.activity, str(section).strip(), str(plant).strip(),
                           q, factor.unit, month_end(y, m), factor)


def seed_scope3(session, registry):
    ws = load_workbook(DATA_XLSX, read_only=True, data_only=True)["Scope 3"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        month, category, description, unit, qty, ef_tco2, _emissions, _scope_type, _mode, _vendor, _stage, source, _method, _unc, _quarter = row[:15]
        if not description or qty is None:
            continue
        y, m = int(str(month)[:4]), int(str(month)[5:7])
        versions = registry.get(3, str(description).strip(), str(unit).strip(), float(ef_tco2), str(source).strip())
        factor = versions[YEAR]
        add_record(session, 3, factor.activity, str(category).strip(), "Central Steel Plant",
                   float(qty), factor.unit, month_end(y, m), factor)
        hist_factor = versions[HIST_YEAR]
        add_record(session, 3, hist_factor.activity, str(category).strip(), "Central Steel Plant",
                   float(qty) * jitter(f"s3|{description}|{y}", 0.85, 0.95), hist_factor.unit,
                   month_end(HIST_YEAR, m), hist_factor)


def seed_business_metrics(session):
    for y in (YEAR, HIST_YEAR):
        for m in range(1, 13):
            session.add(BusinessMetric(
                date=month_end(y, m), metric_name="Tons of Steel Produced",
                value=round(jitter(f"steel|{y}|{m}", 42000.0, 56000.0), 1), unit="tonnes",
            ))
            session.add(BusinessMetric(
                date=month_end(y, m), metric_name="Number of Employees",
                value=float(int(jitter(f"emp|{y}|{m}", 1150.0, 1280.0))), unit="employees",
            ))


def seed_demo_override(session):
    record = (
        session.query(EmissionRecord)
        .filter(EmissionRecord.scope == 1, EmissionRecord.activity_date >= date(YEAR, 1, 1))
        .order_by(EmissionRecord.emissions_kgco2e.desc())
        .first()
    )
    if record is None:
        return
    new_value = round(record.emissions_kgco2e * 0.92, 2)
    record.overridden = True
    record.overridden_value_kgco2e = new_value
    session.add(AuditLog(
        record_id=record.id, action="OVERRIDE",
        old_value_kgco2e=record.emissions_kgco2e, new_value_kgco2e=new_value,
        reason="Sample override: recalculated after fuel meter calibration audit",
        changed_by="sustainability.admin",
    ))


def is_seeded() -> bool:
    with SessionLocal() as session:
        return session.query(EmissionFactor).count() > 0


def run_seed(force: bool = False) -> None:
    Base.metadata.create_all(engine)
    if is_seeded() and not force:
        return
    if force:
        Base.metadata.drop_all(engine)
        Base.metadata.create_all(engine)
    with SessionLocal() as session:
        registry = FactorRegistry(session)
        seed_scope1(session, registry)
        seed_scope2(session, registry)
        seed_scope3(session, registry)
        seed_business_metrics(session)
        session.commit()
        seed_demo_override(session)
        session.commit()


if __name__ == "__main__":
    run_seed()
    with SessionLocal() as s:
        print("factors:", s.query(EmissionFactor).count())
        print("records:", s.query(EmissionRecord).count())
        print("metrics:", s.query(BusinessMetric).count())
        print("audit:", s.query(AuditLog).count())
